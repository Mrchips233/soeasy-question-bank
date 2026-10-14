"""Public offline smoke test using only original generated test artwork.

This module never captures or clicks the desktop. All questions and images are
original test data; bundled PNGs are generated from common color facts and do
not include third-party software screenshots. Rendering fonts are not bundled.
"""
from __future__ import annotations

import importlib
import json
import re
import sys
import time
import traceback
from dataclasses import asdict, replace
from datetime import datetime
from pathlib import Path

EXPECTED_STEM = "蓝色和黄色混合后得到什么颜色？"
EXPECTED_OPTIONS = {"A": "红色", "B": "绿色", "C": "紫色", "D": "黑色"}

def _fixture(filename: str) -> Path:
    base = Path(__file__).resolve().parent
    candidate = base / "_selftest" / filename
    if not candidate.is_file() and not getattr(sys, "frozen", False):
        candidate = base / "public_test_assets" / filename
    if not candidate.is_file():
        raise FileNotFoundError(f"Missing original self-test image: {filename}")
    return candidate

def _check(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def _check_export(output: Path, question, frame, answer_frame, labels: list[str],
                  source: str, simulated: bool) -> dict:
    from docx import Document
    from question_bank import BankStore, save_frame

    store = BankStore(output, question_type=question.question_type)
    question_path, answer_path = store.evidence_paths(question.number, overwrite=True)
    save_frame(frame, question_path)
    save_frame(answer_frame, answer_path)
    _check(store.add(question, labels, question_path, answer_path,
                     answer_source=source, overwrite=True), "Record was not saved")
    saved = json.loads(store.path.read_text(encoding="utf-8"))
    record = next(item for item in saved["questions"] if item["number"] == question.number)
    _check(record["question"] == question.stem, "Saved JSON changed the question")
    _check(record["options"] == question.options, "Saved JSON changed the options")
    _check(record["correct_labels"] == labels, "Saved JSON changed the labels")
    _check(record["correct_answers"] == [question.options[label] for label in labels],
           "Saved JSON changed the answer text")
    _check(record["answer_source"] == source, "Saved JSON changed the inference source")
    _check(saved["total"] == question.total and saved["bank_name"] == question.bank_name,
           "Saved JSON changed the bank identity")
    docx_path = store.export(title="原创离线软件自测（测试答案）")
    _check(docx_path is not None and docx_path.is_file(), "Word export is missing")
    paragraphs = [paragraph.text for paragraph in Document(docx_path).paragraphs]
    expected = [f"{question.number}. {question.stem}"]
    expected.extend(f"{label}. {text}" for label, text in question.options.items())
    expected.append("正确答案：" + "；".join(
        f"{label}. {question.options[label]}" for label in labels))
    _check(all(text in paragraphs for text in expected), "Word export lost a required field")
    return {"number": question.number, "total": question.total,
            "question": question.stem, "options": question.options,
            "tested_labels": labels, "answer_source": source,
            "answer_is_simulated": simulated, "json": str(store.path),
            "word": str(docx_path), "word_required_fields_verified": True}


def _check_gui(output: Path) -> dict:
    import tkinter as tk
    from PIL import Image, ImageTk
    import app
    from question_bank import BASE

    expected_base = (Path(sys.executable).resolve().parent
                     if getattr(sys, "frozen", False)
                     else Path(__import__("question_bank").__file__).resolve().parent)
    _check(BASE.resolve() == expected_base, "Writable BASE is not the expected program directory")
    _check(app.BASE.resolve() == BASE.resolve(), "App and collector disagree about BASE")
    # Exercise App's normal default-path logic without opening any existing user config.
    isolated_base = output / "gui_isolated"
    isolated_base.mkdir(parents=True, exist_ok=True)
    _check(not (isolated_base / "config.json").exists(), "GUI test directory contains a config")
    original_base = app.BASE
    root = None
    try:
        app.BASE = isolated_base
        root = tk.Tk()
        root.withdraw()
        application = app.App(root)
        root.update_idletasks()
        _check(application.config_path == isolated_base / "config.json", "Wrong config path")
        _check(Path(application.output_var.get()) == isolated_base / "output", "Wrong output path")
        _check(len(application.controls) == 14 and all(
            widget.winfo_exists() for widget in application.controls), "GUI controls are missing")
        _check(len(application.question_type_buttons) == 3 and
               {str(button.cget("value")) for button in application.question_type_buttons}
               == {"single", "multiple", "true_false"}, "Question-type controls are missing")
        application.question_type_var.set("true_false")
        _check(application.count_var.get() == "2"
               and str(application.count_entry["state"]) == "disabled"
               and application.current_output_dir() == isolated_base / "output" / "判断题",
               "Judgment-mode GUI did not select two options and its own directory")
        application.question_type_var.set("multiple")
        _check(application.count_var.get() == "4"
               and str(application.count_entry["state"]) == "readonly"
               and application.current_output_dir() == isolated_base / "output" / "多选题",
               "Multiple-choice GUI did not restore the choice count and directory")
        application.question_type_var.set("single")
        _check(application.current_output_dir() == isolated_base / "output",
               "Single-choice GUI did not retain the original output directory")
        _check(application.ocr_btn.winfo_exists(), "Live OCR toggle is missing")
        _check(application.live_ocr.top.state() == "withdrawn", "Hidden OCR monitor opened unexpectedly")
        application.on_ocr({"stage": "preview", "number": 1, "attempt": 1,
                            "text": "1. 原创 OCR 窗口测试\nA. 测试选项", "frame_id": 1})
        application.pump()
        _check(application.live_ocr.latest.get("number") == 1,
               "Live OCR queue did not update the monitor")
        _check("测试选项" in application.live_ocr.text.get("1.0", "end-1c")
               and str(application.live_ocr.text["state"]) == "disabled",
               "Live OCR text is missing or editable")
        photo = ImageTk.PhotoImage(Image.new("RGB", (2, 2), "white"), master=root)
        _check(photo.width() == 2 and photo.height() == 2, "Pillow/Tk image bridge failed")
        _check(not application.config_path.exists() and not (isolated_base / "output").exists(),
               "GUI test unexpectedly wrote config/output")
        return {"hidden_app_created": True, "controls": len(application.controls),
                "question_type_controls_verified": True, "type_switching_verified": True,
                "live_ocr_queue_verified": True, "live_ocr_read_only": True,
                "tcl_version": root.tk.call("info", "patchlevel"),
                "program_base": str(BASE), "normal_config_path": str(BASE / "config.json"),
                "normal_default_output": str(BASE / "output"),
                "test_base": str(isolated_base), "user_data_written": False}
    finally:
        app.BASE = original_base
        if root is not None:
            root.destroy()


def _check_current_collector_features(output: Path, source_frame) -> dict:
    """Exercise current collector behavior using in-memory synthetic desktop pages."""
    import question_bank as backend
    from question_bank import BankStore, Collector, Frame, OCRLine, find_button, parse_question

    _check(backend.REVEAL_CONFIRMATION_ROUNDS == 2, "Reveal confirmation count is stale")
    _check(backend.REVEAL_OBSERVATION_SECONDS == .45, "Reveal observation duration is stale")
    _check(backend.REVEAL_SAMPLE_INTERVAL_SECONDS == .025, "Reveal sampling interval is stale")
    _check(backend.NEXT_MAX_CLICKS == 4, "Next-click retry limit is stale")
    question = parse_question(source_frame)

    def clone(frame, number=None):
        image = frame.image.copy()
        lines = [replace(line) for line in frame.lines]
        # OCR may merge several bottom buttons into one OCR row. These pages
        # are synthetic and need explicit synthetic controls, rather than running
        # additional real OCR on cropped pixels in this behavioral self-test.
        lines = [row for row in lines if row.center[1] < image.height * .91]
        lines.extend([
            OCRLine("查看答案", .99, image.width * .6, image.height * .94,
                    image.width * .76, image.height * .98),
            OCRLine("下一题", .99, image.width * .8, image.height * .94,
                    image.width * .98, image.height * .98),
        ])
        if number is not None:
            old = str(question.number)
            for line in lines:
                line.text = re.sub(r"^" + old + r"(?=[·．.、。?？]|\s*/)",
                                   str(number), line.text)
            # A real pixel difference lets retry logic distinguish the newer page.
            image.putpixel((0, 0), (number % 256, 31, 79))
        return Frame(image, lines)

    class SyntheticStop:
        def __init__(self):
            self.pauses = []

        def check(self):
            pass

        def pause(self, seconds):
            self.pauses.append(seconds)

    class SyntheticDesktop:
        def __init__(self, drop_first_next=False):
            self.pages = [clone(source_frame), clone(source_frame, question.number + 1)]
            self.index = 0
            self.actions = []
            self.next_clicks = 0
            self.drop_first_next = drop_first_next
            self.body_drift = False

        def screenshot(self):
            frame = clone(self.pages[self.index])
            if self.body_drift and self.index == 0:
                for row in frame.lines:
                    if row.text.startswith("D."):
                        row.text = "D.离线模拟后续识别变化"
            frame.image.selftest_frame = frame
            return frame.image

        def click(self, point):
            frame = self.pages[self.index]
            if point == find_button(frame, "下一题").center:
                self.next_clicks += 1
                self.actions.append("下一题")
                if not self.drop_first_next or self.next_clicks > 1:
                    self.index = 1
            elif point == find_button(frame, "查看答案").center:
                self.actions.append("查看答案")
                self.body_drift = True
            else:
                raise AssertionError(f"Unexpected simulated click: {point}")

    class SyntheticEngine:
        def __init__(self):
            self.reads = 0

        def read(self, image, scale=None):
            self.reads += 1
            frame = image.selftest_frame
            return Frame(image, [replace(line) for line in frame.lines])

    updates = []
    desktop, engine, stop = SyntheticDesktop(), SyntheticEngine(), SyntheticStop()
    worker = Collector(desktop, engine, BankStore(output / "direct"), stop,
                       log=lambda message: None, on_ocr=updates.append)
    current, initial = worker.stable_question()
    _check(engine.reads == 1 and current.number == question.number,
           "A complete first question unnecessarily required repeated OCR")
    _check(len(updates) == 1 and updates[0]["stage"] == "stable_question",
           "Collector did not publish its OCR frame")
    snapshot = updates[0]["text"]
    initial.lines[0].text = "Changed after publishing"
    _check(updates[0]["text"] == snapshot, "Live OCR payload retained mutable frame text")

    # Popup body may be entirely hidden; original options still map the explicit answer.
    def forbidden(*args, **kwargs):
        raise AssertionError("Removed body recovery/confirmation was called")

    worker.recover_visible_answer_frame = forbidden
    worker.recover_answer_frame = forbidden
    answer_frame = clone(source_frame)
    answer_frame.lines = [row for row in answer_frame.lines
                          if re.fullmatch(r"\s*\d+\s*/\s*\d+\s*", row.text)
                          or "共" in row.text and "题" in row.text]
    answer_frame.lines.append(OCRLine("正确答案是：C", .99, 100, 500, 400, 540))
    labels, evidence = worker.read_verified_answer(question, answer_frame, "A")
    _check(labels == ["C"] and evidence is answer_frame,
           "Explicit answer without visible body failed to match original options")

    worker.question_frame = clone(source_frame)
    worker.last_frame = worker.question_frame
    worker.reveal_answer(question, worker.question_frame)
    inferred, evidence, dismiss = worker.wait_answer(question, "A")
    _check(inferred == ["A"] and not dismiss
           and desktop.actions == ["查看答案", "查看答案"],
           "No-error inference did not retain exactly two reveal clicks")
    _check(getattr(evidence, "reveal_confirmation_rounds", None) == 2
           and stop.pauses.count(.45) == 2,
           "No-error inference did not complete both reveal observations")
    _check(question.options == current.options,
           "Later body OCR changed the initial question options")

    navigation = SyntheticDesktop(drop_first_next=True)
    navigation_updates = []
    navigator = Collector(navigation, SyntheticEngine(), BankStore(output / "navigation"),
                          SyntheticStop(), log=lambda message: None,
                          on_ocr=navigation_updates.append)
    navigator.question_frame = clone(source_frame)
    next_question, _ = navigator.next_question(question, clone(source_frame))
    _check(next_question.number == question.number + 1
           and navigation.actions == ["下一题", "下一题"],
           "Dropped first Next click was not recovered without skipping a page")
    _check(any(update["stage"] == "next_question" for update in navigation_updates),
           "Next-page OCR did not reach the live viewer callback")
    return {"reveal_rounds": backend.REVEAL_CONFIRMATION_ROUNDS,
            "reveal_observation_seconds": backend.REVEAL_OBSERVATION_SECONDS,
            "reveal_sample_interval_seconds": backend.REVEAL_SAMPLE_INTERVAL_SECONDS,
            "next_max_clicks": backend.NEXT_MAX_CLICKS,
            "single_complete_question_read": True, "immutable_ocr_callback": True,
            "explicit_answer_without_body_recovery": True,
            "two_reveal_clicks_verified": True,
            "next_dropped_click_recovered": True,
            "next_clicks": navigation.next_clicks,
            "next_question": next_question.number,
            "desktop_captured": False, "desktop_clicked": False,
            "answers_are_simulated": True}


def _check_question_type_modes(output: Path) -> dict:
    """Verify all three formats and isolated stores; every page and answer is simulated."""
    from PIL import Image
    from docx import Document
    from question_bank import (BankStore, Frame, OCRLine, parse_answer, parse_question,
                               question_output_dir, save_frame)

    type_names = {"single": "单选题", "true_false": "判断题", "multiple": "多选题"}
    filenames = {"single": "单选题库.docx", "true_false": "判断题库.docx", "multiple": "多选题库.docx"}
    choice_options = {"A": "红色", "B": "绿色", "C": "紫色", "D": "黑色"}

    def page(mode, popup=None):
        image = Image.new("RGB", (670, 1160), "white")
        rows = [OCRLine(f"{type_names[mode]}练习(共7题)", .99, 30, 25, 335, 65),
                OCRLine("1·离线软件自测题目，以下内容只用于验证程序。", .99, 38, 130, 620, 165),
                OCRLine("1/7", .99, 550, 270, 635, 305)]
        options = {"对": "对", "错": "错"} if mode == "true_false" else choice_options
        if popup is None:
            rows.extend(OCRLine(text if mode == "true_false" else f"{label}.{text}",
                                .99, 112, 360 + index * 100, 350, 398 + index * 100)
                        for index, (label, text) in enumerate(options.items()))
        else:
            rows.extend([OCRLine("答题错误", .99, 280, 425, 381, 464),
                         OCRLine(f"正确答案是：{popup}", .99, 210, 480, 455, 520),
                         OCRLine("知道了", .99, 292, 565, 367, 600)])
        rows.extend([OCRLine("查看答案", .99, 408, 1090, 518, 1130),
                     OCRLine("下一题", .99, 550, 1090, 637, 1130)])
        return Frame(image, sorted(rows, key=lambda row: (row.center[1], row.x1)))

    modes, snapshots = {}, []
    for mode, popup, expected in (("single", "C", ["C"]),
                                  ("true_false", "错", ["错"]),
                                  ("multiple", "AB", ["A", "B"])):
        frame, answer_frame = page(mode), page(mode, popup)
        question = parse_question(frame, question_type=mode)
        _check(question.question_type == mode and question.number == 1,
               f"{mode} question type or number changed")
        labels = parse_answer(answer_frame, question)
        _check(labels == expected, f"{mode} explicit answer did not map all original labels")
        if mode == "true_false":
            _check(question.options == {"对": "对", "错": "错"}, "Judgment labels changed")
        store = BankStore(output, question_type=mode)
        _check(store.output_dir == question_output_dir(output, mode), "Mode-specific store path changed")
        question_path, answer_path = store.evidence_paths(question.number)
        save_frame(frame, question_path)
        save_frame(answer_frame, answer_path)
        _check(store.add(question, labels, question_path, answer_path), "Synthetic mode record was not saved")
        saved = json.loads(store.path.read_text(encoding="utf-8"))
        record = saved["questions"][0]
        _check(saved["question_type"] == mode and record["question_type"] == mode,
               "JSON lost its question type")
        _check(record["correct_labels"] == expected
               and record["correct_answers"] == [question.options[label] for label in expected],
               "JSON lost a mode-specific correct answer")
        word = store.export(title="离线软件自测（所有题目和答案均为模拟）")
        _check(word is not None and word.name == filenames[mode], "Mode-specific Word filename changed")
        paragraphs = [paragraph.text for paragraph in Document(word).paragraphs]
        expected_answer = "正确答案：" + "；".join(
            question.options[label] if mode == "true_false" else f"{label}. {question.options[label]}"
            for label in expected)
        expected_options = [text if mode == "true_false" else f"{label}. {text}"
                            for label, text in question.options.items()]
        _check(f"1. {question.stem}" in paragraphs and expected_answer in paragraphs
               and all(text in paragraphs for text in expected_options),
               "Mode-specific Word lost the question, options or complete answer")
        _check(BankStore(output, question_type=mode).records[1]["correct_labels"] == expected,
               "Mode-specific JSON did not resume correctly")
        snapshots.extend([(store.path, store.path.read_bytes()), (word, word.read_bytes())])
        modes[mode] = {"number": 1, "labels": expected, "json": str(store.path),
                       "word": str(word), "word_required_fields_verified": True,
                       "answer_is_simulated": True}
    _check(len({path for path, content in snapshots}) == 6,
           "Modes share a JSON or Word destination")
    _check(all(path.read_bytes() == content for path, content in snapshots),
           "Writing one type overwrote an earlier JSON or Word")
    return {"three_types_verified": True, "separate_json_word_files": True,
            "same_question_number_verified": True, "json_resume_verified": True,
            "answers_are_simulated": True, "desktop_captured": False,
            "desktop_clicked": False, "modes": modes}


def _check_real_synthetic_ocr(output: Path, engine) -> tuple[dict, object]:
    """Run the real local OCR/parser/export stack on wholly original PNG artwork."""
    from PIL import Image
    from question_bank import ocr_text_identity, parse_answer, parse_question

    with Image.open(_fixture("public_question.png")) as image:
        question_frame = engine.read(image.convert("RGB"))
    question = parse_question(question_frame, question_type="single")
    _check((question.number, question.total) == (1, 7), "Original OCR question identity mismatch")
    _check(ocr_text_identity(question.stem) == ocr_text_identity(EXPECTED_STEM),
           "Original OCR question stem mismatch")
    _check(question.options == EXPECTED_OPTIONS, "Original OCR option boundaries mismatch")
    with Image.open(_fixture("public_answer.png")) as image:
        answer_frame = engine.read(image.convert("RGB"))
    labels = parse_answer(answer_frame, question, selected_label="A",
                          reference_frame=question_frame)
    _check(labels == ["B"], "Original OCR explicit answer must be B")
    exported = _check_export(output, question, question_frame, answer_frame,
                             labels, "answer_dialog", simulated=True)
    _check(Path(exported["word"]).name == "单选题库.docx", "Original OCR Word filename mismatch")
    return ({**exported, "correct_labels": labels, "real_local_ocr_verified": True,
             "fixture_origin": "Original color-mixing sample and original Pillow artwork",
             "question_ocr": [asdict(line) for line in question_frame.lines],
             "answer_ocr": [asdict(line) for line in answer_frame.lines],
             "desktop_captured": False, "desktop_clicked": False}, question_frame)


def main(output_dir: str | Path) -> int:
    """Write result.json and return 0/1; no desktop or user data is touched."""
    output = Path(output_dir).expanduser().resolve()
    started = time.perf_counter()
    result = {"status": "running", "success": False,
              "started_at": datetime.now().astimezone().isoformat(),
              "frozen": bool(getattr(sys, "frozen", False)),
              "test_scope": "original PNG artwork, real local OCR, synthetic collector, JSON/Word, hidden GUI",
              "fixtures_are_original": True, "answers_are_test_data": True,
              "network_used": False, "desktop_captured": False,
              "desktop_clicked": False, "checks": {}}
    stage = "create_output"
    try:
        output.mkdir(parents=True, exist_ok=True)
        stage = "dynamic_dependencies"
        versions = {}
        for name in ("pyautogui", "mss", "PIL.ImageTk", "cv2", "onnxruntime", "numpy", "docx"):
            module = importlib.import_module(name)
            versions[name] = str(getattr(module, "__version__", "imported"))
        result["checks"][stage] = versions

        stage = "hidden_gui"
        result["checks"][stage] = _check_gui(output)

        stage = "question_type_modes"
        result["checks"][stage] = _check_question_type_modes(output / "question_types")

        stage = "load_local_ocr"
        from question_bank import OCREngine
        engine = OCREngine()
        result["checks"][stage] = engine.engine.metadata

        stage = "real_synthetic_ocr"
        checked, original_frame = _check_real_synthetic_ocr(output / "original_ocr", engine)
        result["checks"][stage] = checked

        stage = "current_collector_features"
        result["checks"][stage] = _check_current_collector_features(
            output / "current_features", original_frame)

        result.update(status="passed", success=True)
        code = 0
    except Exception as exc:
        result.update(status="failed", failed_stage=stage, error=str(exc),
                      exception_type=type(exc).__name__, traceback=traceback.format_exc())
        code = 1
    result["elapsed_seconds"] = round(time.perf_counter() - started, 3)
    try:
        output.mkdir(parents=True, exist_ok=True)
        temporary = output / "result.json.tmp"
        temporary.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(output / "result.json")
    except OSError:
        return 1
    return code


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1]))
